"""Turning three kinds of source into chunks the rest of the system can use.

Textbook chunking is structural: headings first, then size, never mid-table or
mid-formula. Slide and note handling differ because a slide deck is not prose
and a photograph of handwriting is not text at all.
"""

from __future__ import annotations

import hashlib
import re
from collections.abc import Iterator, Sequence
from pathlib import Path

from .config import Settings
from .llm import LlmClient, PromptLibrary
from .models import ChapterRange, Chunk, ChunkType, NotePage, Parent, SlidePage
from .store import Places


class IngestError(RuntimeError):
    pass


class OutlineMissing(IngestError):
    """The textbook PDF has no usable outline; the UI must supply page ranges."""


_HEADING = re.compile(r"^(#{1,4})\s+(.*)$", re.MULTILINE)
_TABLE_ROW = re.compile(r"^\s*\|", re.MULTILINE)
_FORMULA_FENCE = re.compile(r"\$\$")


def estimate_tokens(text: str) -> int:
    """Four characters per token. Close enough to enforce a budget against."""
    return max(1, len(text) // 4)


def content_hash(data: bytes) -> str:
    return hashlib.blake2b(data, digest_size=16).hexdigest()


def chapter_ranges(pdf: Path) -> list[ChapterRange]:
    """Read chapter boundaries from the PDF outline.

    Raises rather than guessing. A wrong page range silently poisons chapter
    scoping (spec 7.2), so the UI override exists for exactly this case.
    """
    import pypdf
    from pypdf.generic import Destination

    reader = pypdf.PdfReader(str(pdf))
    outline = reader.outline
    entries: list[tuple[str, int]] = []

    def walk(items: object) -> Iterator[tuple[str, int]]:
        if isinstance(items, list):
            for item in items:
                yield from walk(item)
        elif isinstance(items, Destination):
            page = reader.get_destination_page_number(items)
            if page is not None:
                yield str(items.title), int(page) + 1

    entries = list(walk(outline))
    if not entries:
        raise OutlineMissing(f"{pdf.name} has no outline; set page ranges manually")

    total = len(reader.pages)
    ranges: list[ChapterRange] = []
    for index, (title, start) in enumerate(entries):
        end = entries[index + 1][1] - 1 if index + 1 < len(entries) else total
        if end < start:
            end = start
        ranges.append(
            ChapterRange(
                chapter=_slug(title), title=title.strip(), page_start=start, page_end=end
            )
        )
    return ranges


def _slug(text: str) -> str:
    cleaned = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return cleaned or "chapter"


class TextbookIngestor:
    """Structural chunking into parent sections and the child chunks that index them."""

    def __init__(self, settings: Settings) -> None:
        self._chunks = settings.chunks

    def ingest(
        self, pdf: Path, course: str, ranges: Sequence[ChapterRange]
    ) -> tuple[list[Parent], list[Chunk]]:
        pages = self._pages(pdf)
        parents: list[Parent] = []
        chunks: list[Chunk] = []
        for span in ranges:
            body = "\n".join(pages[span.page_start - 1 : span.page_end])
            for parent in self._sections(body, course, span):
                parents.append(parent)
                chunks.extend(self._children(parent))
        return parents, chunks

    @staticmethod
    def _pages(pdf: Path) -> list[str]:
        import pypdf

        reader = pypdf.PdfReader(str(pdf))
        return [page.extract_text() or "" for page in reader.pages]

    def _sections(self, body: str, course: str, span: ChapterRange) -> list[Parent]:
        """Split on headings first. A section longer than parent_tokens is split
        again at paragraph boundaries, never inside a table or a display formula."""
        marks = list(_HEADING.finditer(body))
        if not marks:
            blocks = [(span.title or span.chapter, body)]
        else:
            blocks = []
            for index, mark in enumerate(marks):
                start = mark.end()
                stop = marks[index + 1].start() if index + 1 < len(marks) else len(body)
                blocks.append((mark.group(2).strip(), body[start:stop]))

        parents: list[Parent] = []
        for heading, text in blocks:
            for part, piece in enumerate(self._split_to_size(text, self._chunks.parent_tokens)):
                if len(piece.strip()) < self._chunks.min_chunk_chars:
                    continue
                identifier = f"{course}:{span.chapter}:{_slug(heading)}:{part}"
                parents.append(
                    Parent(
                        id=identifier,
                        course=course,
                        chapter=span.chapter,
                        heading=heading,
                        section_path=f"{span.title or span.chapter} > {heading}",
                        text=piece.strip(),
                        page_start=span.page_start,
                        page_end=span.page_end,
                        token_estimate=estimate_tokens(piece),
                    )
                )
        return parents

    @staticmethod
    def _split_to_size(text: str, limit: int) -> list[str]:
        if estimate_tokens(text) <= limit:
            return [text]
        pieces: list[str] = []
        current: list[str] = []
        size = 0
        in_formula = False
        for line in text.splitlines(keepends=True):
            if _FORMULA_FENCE.search(line):
                in_formula = not in_formula
            line_tokens = estimate_tokens(line)
            breakable = not in_formula and not _TABLE_ROW.match(line) and not line.strip()
            if size + line_tokens > limit and current and breakable:
                pieces.append("".join(current))
                current, size = [], 0
            current.append(line)
            size += line_tokens
        if current:
            pieces.append("".join(current))
        return pieces

    def _children(self, parent: Parent) -> list[Chunk]:
        pieces = self._split_to_size(parent.text, self._chunks.child_tokens)
        children: list[Chunk] = []
        for index, piece in enumerate(pieces):
            if len(piece.strip()) < self._chunks.min_chunk_chars:
                continue
            children.append(
                Chunk(
                    id=f"{parent.id}#{index}",
                    course=parent.course,
                    chapter=parent.chapter,
                    text=piece.strip(),
                    section_path=parent.section_path,
                    page_start=parent.page_start,
                    page_end=parent.page_end,
                    parent_id=parent.id,
                    chunk_type=ChunkType.TABLE if _TABLE_ROW.search(piece) else ChunkType.PROSE,
                    token_estimate=estimate_tokens(piece),
                )
            )
        return children


class SlideIngestor:
    """One slide is one chunk, with a `## Page N` marker kept for citation."""

    FIGURE_TEXT_FLOOR = 40

    def __init__(self, settings: Settings, client: LlmClient) -> None:
        self._settings = settings
        self._client = client
        self._prompts = PromptLibrary(settings.prompts_dir)

    def ingest(self, source: Path) -> list[SlidePage]:
        if source.is_dir():
            decks = sorted(
                path for path in source.iterdir() if path.suffix.lower() in {".pdf", ".pptx"}
            )
        else:
            decks = [source]
        if not decks:
            raise IngestError(f"No slide decks found at {source}")
        pages: list[SlidePage] = []
        for deck in decks:
            pages.extend(
                self._pptx(deck) if deck.suffix.lower() == ".pptx" else self._pdf(deck)
            )
        return pages

    def _pdf(self, deck: Path) -> list[SlidePage]:
        import pypdf

        reader = pypdf.PdfReader(str(deck))
        return [
            self._page(deck.stem, number, page.extract_text() or "")
            for number, page in enumerate(reader.pages, start=1)
        ]

    def _pptx(self, deck: Path) -> list[SlidePage]:
        from pptx import Presentation

        pages: list[SlidePage] = []
        for number, slide in enumerate(Presentation(str(deck)).slides, start=1):
            body = "\n".join(
                shape.text_frame.text for shape in slide.shapes if shape.has_text_frame
            )
            notes = ""
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text
            if notes.strip():
                body = f"{body}\n\n_Speaker notes:_ {notes.strip()}"
            pages.append(self._page(deck.stem, number, body))
        return pages

    def _page(self, deck: str, number: int, text: str) -> SlidePage:
        sparse = len(text.strip()) < self.FIGURE_TEXT_FLOOR
        return SlidePage(
            page=number,
            deck=deck,
            markdown=f"## Page {number}\n\n{text.strip()}",
            is_figure_only=sparse,
        )

    def describe_figures(self, pages: Sequence[SlidePage], images: Path) -> list[SlidePage]:
        """Diagram-only pages get a vision pass. Pure text extraction throws
        away most of what a slide diagram means."""
        prompt = self._prompts.render("describe_figure")
        described: list[SlidePage] = []
        for page in pages:
            image = images / f"page{page.page:03d}.png"
            if not page.is_figure_only or not image.exists():
                described.append(page)
                continue
            description = self._client.vision(prompt, image, job="describe_figure")
            described.append(
                page.model_copy(
                    update={
                        "figure_description": description,
                        "markdown": f"{page.markdown}\n\n[diagram: {description.strip()}]",
                    }
                )
            )
        return described


class NoteIngestor:
    """Handwriting to Markdown, cached by content hash.

    OCR on handwriting is the least reliable step in the system, which is why
    the graph interrupts for review right after this runs.
    """

    def __init__(self, settings: Settings, client: LlmClient, places: Places) -> None:
        self._client = client
        self._places = places
        self._prompts = PromptLibrary(settings.prompts_dir)

    def ingest(self, source: Path) -> list[NotePage]:
        images = sorted(
            path
            for path in (source.iterdir() if source.is_dir() else [source])
            if path.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp"}
        )
        if not images:
            raise IngestError(f"No note images found at {source}")
        prompt = self._prompts.render("ocr_notes")
        pages: list[NotePage] = []
        for number, image in enumerate(images, start=1):
            digest = content_hash(image.read_bytes())
            cache = self._places.note_cache(digest)
            if cache.exists():
                markdown = cache.read_text(encoding="utf-8")
            else:
                markdown = self._client.vision(prompt, image, job="ocr_notes")
                cache.write_text(markdown, encoding="utf-8")
            pages.append(
                NotePage(
                    page=number,
                    image_path=str(image),
                    content_hash=digest,
                    markdown=markdown,
                )
            )
        return pages
